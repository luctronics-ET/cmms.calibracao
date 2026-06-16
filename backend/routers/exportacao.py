"""Exportação do inventário (CSV/XLSX/PDF) a partir de uma lista de ids ordenada."""
from __future__ import annotations
import csv
import io
from datetime import date
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from backend.db import get_db
from backend.models import Instrumento
from backend.servico import instrumento_para_out
from backend.schemas import ExportRequest, InstrumentoOut

router = APIRouter(prefix="/api/v1", tags=["exportacao"])

# (chave no dict de linha, cabeçalho legível) — fonte única da ordem das colunas
COLUNAS: list[tuple[str, str]] = [
    ("codigo_interno", "Código interno"),
    ("codigo_patrimonial", "Patrimônio"),
    ("serial", "Série"),
    ("equipamento", "Equipamento"),
    ("marca", "Marca"),
    ("modelo", "Modelo"),
    ("familia_nome", "Família"),
    ("tipo_nome", "Tipo"),
    ("grandeza_nome", "Grandeza"),
    ("unidade_simbolo", "Unidade"),
    ("faixa", "Faixa"),
    ("resolucao", "Resolução"),
    ("emp", "EMP"),
    ("disciplina", "Disciplina"),
    ("setor", "Setor"),
    ("localizacao", "Localização"),
    ("status_operacional", "Status operacional"),
    ("status", "Status validade"),
    ("dias_restantes", "Dias restantes"),
    ("data_ultima_calibracao", "Última calibração"),
    ("data_validade", "Validade"),
    ("ciclo_meses", "Ciclo (meses)"),
    ("igp", "IGP"),
    ("classe_prioridade", "Prioridade"),
    ("observacoes", "Observações"),
]

MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _faixa_txt(o: InstrumentoOut) -> str:
    if o.faixa_min is not None or o.faixa_max is not None:
        return f"{o.faixa_min if o.faixa_min is not None else ''} … {o.faixa_max if o.faixa_max is not None else ''}"
    return o.faixa or ""


def _localizacao_txt(o: InstrumentoOut) -> str:
    return " → ".join(p for p in [o.organizacao, o.unidade_org, o.secao, o.bancada] if p)


def _valor(o: InstrumentoOut, chave: str):
    """Valor textual de uma coluna para um instrumento (datas em ISO)."""
    if chave == "faixa":
        return _faixa_txt(o)
    if chave == "localizacao":
        return _localizacao_txt(o)
    v = getattr(o, chave)
    if isinstance(v, date):
        return v.isoformat()
    return "" if v is None else v


def _linhas_completas(instrumentos: list[InstrumentoOut]) -> list[dict]:
    """Uma linha-dict por instrumento, com as chaves de COLUNAS."""
    return [{chave: _valor(o, chave) for chave, _ in COLUNAS} for o in instrumentos]


def _gerar_csv(linhas: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([rotulo for _, rotulo in COLUNAS])
    for linha in linhas:
        w.writerow([linha[chave] for chave, _ in COLUNAS])
    # BOM para o Excel pt-BR abrir UTF-8 corretamente
    return buf.getvalue().encode("utf-8-sig")


def _gerar_xlsx(linhas: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventário"
    rotulos = [rotulo for _, rotulo in COLUNAS]
    ws.append(rotulos)
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        ws.append([linha[chave] for chave, _ in COLUNAS])
    # auto-largura simples: maior conteúdo da coluna (limitado a 50)
    for idx, (chave, rotulo) in enumerate(COLUNAS, start=1):
        largura = max([len(rotulo)] + [len(str(l[chave])) for l in linhas]) + 2
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(largura, 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _latin1(texto: str) -> str:
    """Coage texto ao Latin-1 (fonte core Helvetica) substituindo caracteres fora da faixa."""
    return texto.encode("latin-1", "replace").decode("latin-1")


def _gerar_pdf(instrumentos: list[InstrumentoOut]) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos
    hoje = date.today()
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, "CMASM / DME - Inventario de Calibracao",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 6, f"Gerado em {hoje.isoformat()} - {len(instrumentos)} item(ns)",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)

    # cabeçalhos e larguras (mm) do resumo
    cols = [
        ("Codigo", 45), ("Equipamento", 75), ("Setor", 45),
        ("Validade", 30), ("Status", 40), ("IGP/Classe", 42),
    ]
    pdf.set_font("Helvetica", "B", 9)
    for titulo, larg in cols:
        pdf.cell(larg, 7, titulo, border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 8)
    for o in instrumentos:
        validade = o.data_validade.isoformat() if o.data_validade else "-"
        igp = f"{o.igp} / {o.classe_prioridade}" if o.igp is not None else o.classe_prioridade
        valores = [
            o.codigo_interno or o.codigo_patrimonial or "-",
            o.equipamento or "-", o.setor or "-",
            validade, o.status, igp,
        ]
        for (titulo, larg), valor in zip(cols, valores):
            texto = _latin1(str(valor))
            # trunca para caber na célula (resumo)
            while pdf.get_string_width(texto) > larg - 2 and len(texto) > 1:
                texto = texto[:-1]
            pdf.cell(larg, 6, texto, border=1)
        pdf.ln()
    saida = pdf.output()
    return bytes(saida)


@router.post("/instrumentos/export")
def exportar(req: ExportRequest, db: Session = Depends(get_db)):
    if not req.ids:
        raise HTTPException(status_code=400, detail="Nada a exportar")
    hoje = date.today()
    por_id = {i.id: i for i in db.query(Instrumento).filter(Instrumento.id.in_(req.ids)).all()}
    ordenados = [instrumento_para_out(por_id[i], hoje) for i in req.ids if i in por_id]
    linhas = _linhas_completas(ordenados)

    if req.formato == "csv":
        conteudo = _gerar_csv(linhas)
    elif req.formato == "xlsx":
        conteudo = _gerar_xlsx(linhas)
    else:
        conteudo = _gerar_pdf(ordenados)

    nome = f"inventario_{hoje.isoformat()}.{req.formato}"
    return Response(
        content=conteudo,
        media_type=MEDIA[req.formato],
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )
